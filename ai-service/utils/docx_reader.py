"""
utils/docx_reader.py
=====================
Purpose:
    Extracts raw text content from DOCX resume files, including
    paragraph text and text inside tables (many resume templates use
    tables for layout).

Architecture:
    Single public function `extract_text_from_docx` used by
    services/resume_service.py.

Usage Example:
    from utils.docx_reader import extract_text_from_docx
    text = extract_text_from_docx(file_bytes)
"""

from __future__ import annotations

import io

from docx import Document

from utils.exceptions import ResumeParsingError
from utils.logger import logger


def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extract all readable text from a DOCX file's raw bytes, covering
    both paragraphs and table cells.

    Parameters:
        file_bytes: Raw bytes of the uploaded DOCX file.

    Returns:
        The concatenated text content of the document.

    Raises:
        ResumeParsingError: If the document cannot be opened or parsed.
    """
    try:
        document = Document(io.BytesIO(file_bytes))
    except Exception as exc:  # noqa: BLE001
        logger.error(f"Failed to open DOCX file: {exc}")
        raise ResumeParsingError("The uploaded DOCX file is corrupted or invalid") from exc

    chunks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    chunks.append(cell.text)

    text = "\n".join(chunks)

    if not text.strip():
        raise ResumeParsingError("No readable text found inside the DOCX file")

    return text
